"""v3.2.0 — 'AR & AP Ageing' download pack, laid out like the client's
month-end workbook (Dashboard, AR Ageing, AP Ageing), with live formulas."""
import io
import sys
import types
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[1]))
sys.path.insert(0, str(HERE))
import _fakefrappe as ff  # noqa: E402

OUT = Path("/tmp/ageing_pack_test.xlsx")


def fake_ageing(company=None, as_of=None, party_type="Customer", split_not_due=0, **k):
    if party_type == "Customer":
        rows = [{"party": "C1", "label": "Acme Trading", "invoiced": 1000.0, "paid": 300.0, "outstanding": 700.0,
                 "not_due": 0.0, "buckets": [400.0, 200.0, 0.0, 100.0, 0.0]},
                {"party": "C2", "label": "Beta Co", "invoiced": 500.0, "paid": 0.0, "outstanding": 500.0,
                 "not_due": 0.0, "buckets": [0.0, 0.0, 0.0, 0.0, 500.0]}]
    else:
        nd = 150.0 if split_not_due else 0.0
        rows = [{"party": "S1", "label": "Gulf Supplies", "invoiced": 800.0, "paid": 200.0, "outstanding": 600.0,
                 "not_due": nd, "buckets": [450.0 - nd + 150.0 - 150.0 if split_not_due else 450.0, 150.0, 0, 0, 0]}]
    # v3.2.1 — the REAL label builder: it returns {"key","en","ar"} objects, which
    # the first version of the pack wrote straight into Excel (server error).
    from neotec_insight.neotec_insight.api.ageing import _slab_labels
    return {"rows": rows, "labels": _slab_labels([30, 60, 90, 120], "days"), "currency": "SAR"}


class Pack(unittest.TestCase):
    def setUp(self):
        self.site = ff.install()
        import frappe
        self.frappe = frappe
        self.site.table("Company")["Alpha Co"] = {"name": "Alpha Co", "company_name": "شركة ألفا", "default_currency": "SAR"}
        frappe.local.response = types.SimpleNamespace()
        from neotec_insight.neotec_insight.api import ageing, packs
        self.packs = packs
        self.real_ageing = ageing.ar_ap_ageing
        ageing.ar_ap_ageing = fake_ageing
        packs._notes_by_party = lambda c, d, pt: {"C1": 50.0} if pt == "Customer" else {"S1": 20.0}
        packs._seed_ageing_pack()

    def tearDown(self):
        ff.uninstall()

    def build(self):
        self.packs.generate_pack("ar_ap_ageing", "Alpha Co", "2026-09-01", "2026-09-30")
        data = self.frappe.local.response.filecontent
        OUT.write_bytes(data)
        import openpyxl
        return openpyxl.load_workbook(io.BytesIO(data))

    def test_layout_matches_the_template(self):
        wb = self.build()
        self.assertEqual(wb.sheetnames, ["Dashboard", "AR Ageing", "AP Ageing"])
        ar, ap = wb["AR Ageing"], wb["AP Ageing"]
        self.assertEqual([c.value for c in ar[9]][:6], ["#", "Party Name", "Invoiced Amount", "Paid Amount",
                                                        "Credit Note", "Outstanding"])
        self.assertEqual(ar["G9"].value, "Not Yet Due")                       # v3.3.1 — AR too
        self.assertEqual(ar["M9"].value, "Total Due")
        self.assertEqual(ap["E9"].value, "Debit Note")
        self.assertEqual(ap["G9"].value, "Not Yet Due")
        self.assertEqual(ap["M9"].value, "Total Due")
        self.assertIn("30 September 2026", ar["B5"].value)
        self.assertEqual(ar.freeze_panes, "C10")

    def test_notes_split_out_of_payments(self):
        ar = self.build()["AR Ageing"]
        self.assertEqual((ar["D10"].value, ar["E10"].value), (250.0, 50.0))   # paid 300 = 250 payments + 50 notes
        self.assertEqual(ar["F10"].value, "=C10-D10-E10")
        self.assertEqual(ar["M10"].value, "=SUM(G10:L10)")
        self.assertEqual(ar["A9"].fill.fgColor.rgb[-6:], "57257D")          # template heading colour
        self.assertEqual(ar["C12"].value, "=SUM(C10:C11)")                    # total row

    def test_dashboard_links(self):
        d = self.build()["Dashboard"]
        self.assertEqual(d["E6"].value, "='AR Ageing'!G12")
        self.assertEqual(d["I6"].value, "='AP Ageing'!G11")                   # AP: 1 supplier, total row 11, starts with Not Yet Due
        self.assertTrue(str(d["I7"].value).startswith("='AP Ageing'!H"))

    def test_pack_is_seeded_once(self):
        self.packs._seed_ageing_pack()
        self.assertEqual(len([k for k in self.site.table("Insight Export Pack")]), 1)


if __name__ == "__main__":
    unittest.main()


class FromAgeingScreen(Pack):
    """v3.2.1 — the Ageing screen's Download pack button."""

    def test_overrides_reach_every_ageing_sheet_and_bad_values_are_ignored(self):
        seen = []
        from neotec_insight.neotec_insight.api import ageing

        def spy(**k):
            seen.append(k)
            return fake_ageing(**k)
        ageing.ar_ap_ageing = spy
        self.packs.generate_pack("ar_ap_ageing", "Alpha Co", "2026-10-27", "2026-10-27",
                                 overrides='{"slabs": "30,60,90", "based_on": "posting", "mode": "months", "evil": 1}')
        self.assertEqual({(k["slabs"], k["based_on"], k["mode"]) for k in seen}, {("30,60,90", "posting", "months")})
        seen.clear()
        self.packs.generate_pack("ar_ap_ageing", "Alpha Co", "2026-10-27", "2026-10-27",
                                 overrides='{"based_on": "DROP", "mode": "weeks"}')
        self.assertEqual({(k["based_on"], k["mode"]) for k in seen}, {("due", "days")})

    def test_missing_pack_is_created_on_the_fly(self):
        self.site.table("Insight Export Pack").clear()
        self.packs.generate_pack("ar_ap_ageing", "Alpha Co", "2026-10-27", "2026-10-27")
        self.assertIn("ar_ap_ageing", self.site.table("Insight Export Pack"))
        self.assertTrue(self.frappe.local.response.filecontent)


class Labels(Pack):
    def test_bucket_headers_are_text_in_the_pack_language(self):
        ar = self.build()["AR Ageing"]
        self.assertEqual(ar["H9"].value, "0\u201330 Days")
        self.assertEqual(ar["L9"].value, "121+ Days")


def many_customers(company=None, as_of=None, party_type="Customer", split_not_due=0, **k):
    from neotec_insight.neotec_insight.api.ageing import _slab_labels
    if party_type == "Customer":
        amts = [("Alpha", 100.0), ("Bravo", 900.0), ("Charlie", 300.0), ("Delta", 0.0), ("Echo", 700.0),
                ("Foxtrot", 500.0), ("Golf", 800.0)]
        rows = [{"party": n, "label": n, "invoiced": a, "paid": 0.0, "outstanding": a, "not_due": 0.0,
                 "buckets": [a, 0, 0, 0, 0]} for n, a in amts]
    else:
        rows = [{"party": "S1", "label": "Gulf Supplies", "invoiced": 50.0, "paid": 0.0, "outstanding": 50.0,
                 "not_due": 0.0, "buckets": [50.0, 0, 0, 0, 0]}]
    return {"rows": rows, "labels": _slab_labels([30, 60, 90, 120], "days"), "currency": "SAR"}


class TopAndStyle(unittest.TestCase):
    """v3.3.0 — Top 5 receivables always printed; header and colours configurable.
    Shares Pack's setup and build, not its tests (different data)."""
    build = Pack.build
    tearDown = Pack.tearDown

    def setUp(self):
        Pack.setUp(self)
        from neotec_insight.neotec_insight.api import ageing
        ageing.ar_ap_ageing = many_customers
        self.packs._notes_by_party = lambda c, d, pt: {}

    def dash(self, **opts):
        import json as _j
        rec = self.site.table("Insight Export Pack")["ar_ap_ageing"]
        cfg = _j.loads(rec["config_json"])
        cfg["sheets"][0]["options"] = opts
        rec["config_json"] = _j.dumps(cfg)
        wb = self.build()
        return wb["Dashboard"], wb

    def top_rows(self, d):
        start = next(r for r in range(1, 60) if str(d.cell(row=r, column=2).value or "").strip().startswith("Top"))
        out = []
        for r in range(start + 2, start + 12):
            if d.cell(row=r, column=2).value is None:
                break
            out.append((d.cell(row=r, column=2).value, d.cell(row=r, column=3).value, d.cell(row=r, column=9).value))
        return d.cell(row=start, column=2).value.strip(), out

    def test_top5_ranked_over_all_customers_with_live_links(self):
        d, wb = self.dash()
        head, rows = self.top_rows(d)
        self.assertEqual(head, "Top 5 AR Customers by Overdue Balance")
        self.assertEqual([n for _k, n, _f in rows], ["Bravo", "Golf", "Echo", "Foxtrot", "Charlie"])
        self.assertEqual([amt for _k, _n, amt in rows], [900.0, 800.0, 700.0, 500.0, 300.0])   # values, not formulas
        self.assertEqual(d.cell(row=5, column=2).fill.fgColor.rgb[-6:], "262626")   # template defaults

    def test_from_the_ageing_screen_top5_is_still_printed(self):
        self.packs.generate_pack("ar_ap_ageing", "Alpha Co", "2026-10-27", "2026-10-27",
                                 overrides='{"slabs": "30,60,90,120", "based_on": "due", "mode": "days"}')
        import io as _io
        import openpyxl
        d = openpyxl.load_workbook(_io.BytesIO(self.frappe.local.response.filecontent))["Dashboard"]
        self.assertEqual(len(self.top_rows(d)[1]), 5)

    def test_header_and_colours(self):
        d, wb = self.dash(header_title="IRSAA — Month-end Ageing", header_subtitle="Riyadh HQ",
                          color_accent="#0B5394", color_dark="zzz", top_ap="3")
        self.assertEqual(d["B2"].value, "IRSAA — Month-end Ageing")
        self.assertIn("Riyadh HQ", d["B3"].value)
        self.assertEqual(d.cell(row=5, column=2).fill.fgColor.rgb[-6:], "262626")      # invalid → default
        self.assertEqual(wb["AR Ageing"]["A9"].fill.fgColor.rgb[-6:], "0B5394")         # accent applied
        self.assertEqual(wb["AR Ageing"]["B4"].value, "Riyadh HQ")
        heads = [str(d.cell(row=r, column=2).value or "") for r in range(1, 60)]
        self.assertTrue(any("Top 3 AP Suppliers" in h for h in heads))

    def test_top_ar_can_be_hidden(self):
        d, _wb = self.dash(top_ar="0")
        heads = [str(d.cell(row=r, column=2).value or "") for r in range(1, 60)]
        self.assertFalse(any(h.strip().startswith("Top") for h in heads))      # table headings only


def big_but_not_due(company=None, as_of=None, party_type="Customer", split_not_due=0, **k):
    from neotec_insight.neotec_insight.api.ageing import _slab_labels
    if party_type == "Customer":
        # KAFD-like: the biggest balance, all of it not yet due
        spec = [("Big Current", 2000.0, 2000.0), ("Late A", 600.0, 0.0), ("Late B", 900.0, 100.0),
                ("Late C", 300.0, 0.0), ("Late D", 250.0, 50.0), ("Late E", 150.0, 0.0), ("Late F", 90.0, 0.0)]
        rows = [{"party": n, "label": n, "invoiced": o, "paid": 0.0, "outstanding": o,
                 "not_due": nd if split_not_due else 0.0,
                 "buckets": [o - nd if split_not_due else o, 0, 0, 0, 0]} for n, o, nd in spec]
    else:
        rows = []
    return {"rows": rows, "labels": _slab_labels([30, 60, 90, 120], "days"), "currency": "SAR"}


class OverdueRanking(unittest.TestCase):
    """v3.3.1 — top customers by the amount actually due, not by total balance."""
    build = Pack.build
    tearDown = Pack.tearDown
    top_rows = TopAndStyle.top_rows
    dash = TopAndStyle.dash

    def setUp(self):
        Pack.setUp(self)
        from neotec_insight.neotec_insight.api import ageing
        ageing.ar_ap_ageing = big_but_not_due
        self.packs._notes_by_party = lambda c, d, pt: {}

    def test_not_yet_due_customer_is_not_top(self):
        d, wb = self.dash()
        _h, rows = self.top_rows(d)
        names = [n for _k, n, _a in rows]
        self.assertEqual(names, ["Late B", "Late A", "Late C", "Late D", "Late E"])
        self.assertEqual(rows[0][2], 800.0)                       # 900 outstanding − 100 not yet due
        top = next(r for r in range(1, 60) if str(d.cell(row=r, column=2).value or "").strip().startswith("Top"))
        self.assertIn("Due", d.cell(row=top + 1, column=9).value)
        self.assertEqual(d.cell(row=top + 2, column=7).value, 900.0)          # outstanding shown alongside
        self.assertTrue(wb.calculation.fullCalcOnLoad)

    def test_can_rank_by_outstanding(self):
        d, _wb = self.dash(top_by="outstanding")
        _h, rows = self.top_rows(d)
        self.assertEqual(rows[0][1], "Big Current")
